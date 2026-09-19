"""
End-to-end test of the multi-rater annotation pipeline on simulated raters:
build packet -> fill workbooks -> import -> expert validation.
"""

import csv
import os

import pytest
from openpyxl import load_workbook

from annotation import build_packet as bp
from annotation import import_returns as ir
from annotation.redaction import redact, REDACTED
from analysis import expert_validation as ev

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ---------------------------------------------------------------------------
# Redaction
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("text,gone", [
    ("typical symptoms of BLB such as water-soaked areas", "BLB"),
    ("Rice blast caused by Pyricularia oryzae Cav. was identified", "Pyricularia"),
    ("olive-green colour smut ball", "smut"),
    ("co-infected with RTBV and RTSV showed", "RTBV"),
    ("characteristic root galls of the genus Meloidogyne", "Meloidogyne"),
    ("Gejala serangan blas ditemukan", "blas"),
])
def test_redaction_removes_names(text, gone):
    out, removed = redact(text)
    assert gone.lower() not in out.lower() and REDACTED in out and removed


def test_redaction_keeps_signs_and_vectors():
    text = "diamond-shaped lesions, velvety spore balls and the green leafhopper Nephotettix virescens"
    assert redact(text) == (text, [])


def test_every_case_text_is_free_of_class_names():
    for c in bp.load_cases():
        low = c["text"].lower()
        for name in ("blast", "blight", "smut", "tungro", "grassy stunt", "nematode", "meloidogyne", "xanthomonas"):
            if name == "blight" and ("blight symptom" in low or "blighting" in low or "leaf blight and" in low):
                continue
            assert name not in low, (c["case_id"], name)


# ---------------------------------------------------------------------------
# Packet structure
# ---------------------------------------------------------------------------

def test_codes_are_stable_and_complete():
    a, b = bp.code_to_case(), bp.code_to_case()
    assert a == b
    with open(os.path.join(BASE_DIR, "data", "benchmark_field.csv"), encoding="utf-8-sig") as f:
        ids = {r["case_id"] for r in csv.DictReader(f)}
    assert set(a.values()) == ids


def test_rater_orders_differ():
    cases = bp.load_cases()
    assert [c["code"] for c in bp.rater_order(cases, 0)] != [c["code"] for c in bp.rater_order(cases, 1)]


def test_symptom_labels_cover_all_observation_terms():
    from ricekg import model
    assert set(bp.symptom_labels(bp.load_definitions()).values()) == set(model.ALL_SYMPTOMS)


# ---------------------------------------------------------------------------
# Round trip with simulated raters
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def packet(tmp_path_factory, request):
    out = tmp_path_factory.mktemp("packet")
    mp = pytest.MonkeyPatch()

    def fake_explain(cases, defs):
        for c in cases:
            c["conclusion"], c["explanation"], c["recorded"] = "x", "y", "z"

    mp.setattr(bp, "explain_cases", fake_explain)
    bp.build(str(out), n_raters=3)
    mp.undo()
    return out


def _fill(packet_dir, returned_dir, rater, noisy_codes=()):
    truth = {c["code"]: c for c in bp.load_cases()}
    dx_label = {k: bp.diagnosis_label(l, k) for l, k in bp.DIAGNOSIS_OPTIONS}
    term_label = {t: l for l, t in bp.symptom_labels(bp.load_definitions()).items()}
    wb = load_workbook(os.path.join(packet_dir, f"{rater}_TahapA.xlsx"))
    ws = wb["Tugas A"]
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 1).value
        c = truth[code]
        key = "Other" if c["truth"] == "No_Diagnosis" else c["truth"]
        if code in noisy_codes:
            key = "Undetermined"
        ws.cell(r, 3).value = dx_label[key]
        ws.cell(r, 4).value = "Yakin"
        for j, term in enumerate(c["symptoms"]):
            ws.cell(r, 5 + j).value = term_label[term]
    rv = wb["Review Definisi"]
    rv.cell(2, 4).value = "Perlu revisi"
    rv.cell(2, 5).value = "usulan"
    wb.save(os.path.join(returned_dir, f"{rater}_TahapA.xlsx"))

    wb = load_workbook(os.path.join(packet_dir, f"{rater}_TahapB.xlsx"))
    ws = wb["Tugas B"]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 6).value, ws.cell(r, 7).value, ws.cell(r, 8).value, ws.cell(r, 9).value = "Ya", "4", "3", "5"
    wb.save(os.path.join(returned_dir, f"{rater}_TahapB.xlsx"))


@pytest.fixture(scope="module")
def imported(packet, tmp_path_factory):
    returned = tmp_path_factory.mktemp("returned")
    _fill(packet, returned, "R1")
    _fill(packet, returned, "R2")
    _fill(packet, returned, "R3", noisy_codes={"K01", "K02", "K03"})
    out = tmp_path_factory.mktemp("data")
    info = ir.import_returns(str(returned), str(out))
    return out, info


def test_import_counts(imported):
    out, info = imported
    assert info["raters"] == ["R1", "R2", "R3"]
    assert info["complete_cases"] == 56 and info["diagnoses"] == 168
    assert info["explanation_rows"] == 168 and info["review_rows"] == 3
    with open(os.path.join(out, "annotations_multirater.csv"), encoding="utf-8") as f:
        header = next(csv.reader(f))
    assert header == ["case_id", "annotator_1", "annotator_2", "annotator_3", "notes"]


def test_import_rejects_free_text_in_dropdown(packet, tmp_path):
    _fill(packet, tmp_path, "R1")
    wb = load_workbook(tmp_path / "R1_TahapA.xlsx")
    wb["Tugas A"].cell(2, 3).value = "blas saja"
    wb.save(tmp_path / "R1_TahapA.xlsx")
    with pytest.raises(ir.ImportError_):
        ir.import_returns(str(tmp_path), str(tmp_path / "out"))


def test_expert_validation_on_simulated_raters(imported):
    out, _ = imported
    rep = ev.run(str(out))
    dx = rep["diagnosis"]
    assert dx["n_cases"] == 56
    assert dx["pairwise_rater_kappa"]["R1-R2"] == 1.0
    assert dx["pairwise_rater_kappa"]["R1-R3"] < 1.0
    assert dx["vs_published_label"]["R1"]["all"] == {"k": 56, "n": 56}
    assert rep["encoding"]["author_vs_majority_jaccard"] == 1.0
    assert rep["explanations"]["overall"]["accept"] == {"yes": 168}
    assert list(rep["definitions"]["flagged_terms"]) == [sorted(bp.load_definitions())[0]]


def test_expert_validation_refuses_without_data(tmp_path):
    assert ev.run(str(tmp_path)) is None


@pytest.mark.parametrize("outputs,label", [
    ([("Rice_Blast", "suspected")], "Rice_Blast"),
    ([("Rice_Blast", "suspected"), ("False_Smut", "suspected")], "Multiple"),
    ([("x", "out_of_scope")], "Other"),
    ([("Rice_Blast", "possible")], "Undetermined"),
    ([], "Undetermined"),
])
def test_system_label(outputs, label):
    assert ev.system_label(outputs) == label


# ---------------------------------------------------------------------------
# Re-rating packet after ruleset v2.4.0
# ---------------------------------------------------------------------------

def test_rerating_packet_round_trip(tmp_path, monkeypatch):
    from annotation import build_rerating_packet as rr

    def fake_explain(cases, defs):
        for c in cases:
            c["conclusion"] = "blas (DIDUGA)" if c["code"] == "K01" else "Tidak ada diagnosis"
            c["explanation"], c["recorded"] = "y", "z"

    monkeypatch.setattr(bp, "explain_cases", fake_explain)
    packet_dir, returned = tmp_path / "packet", tmp_path / "returned"
    info = rr.build(str(packet_dir), n_raters=2, review=("R2",))
    assert info["cases"] == 56
    assert sorted(os.listdir(packet_dir)) == ["R1_TahapB_v24.xlsx", "R2_ReviewDefinisi.xlsx",
                                              "R2_TahapB_v24.xlsx", "Surat_R1.txt", "Surat_R2.txt"]

    # Stage A for both raters is required by the importer; reuse the first-round packet builder.
    first = tmp_path / "first"
    monkeypatch.setattr(bp, "explain_cases", lambda cases, defs: [
        c.update(conclusion="x", explanation="y", recorded="z") for c in cases])
    bp.build(str(first), n_raters=2)
    returned.mkdir()
    for rater in ("R1", "R2"):
        _fill(first, returned, rater)
        wb = load_workbook(packet_dir / f"{rater}_TahapB_v24.xlsx")
        ws = wb["Tugas B"]
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 6).value, ws.cell(r, 7).value, ws.cell(r, 8).value, ws.cell(r, 9).value = "Sebagian", "3", "3", "4"
        wb.save(returned / f"{rater}_TahapB_v24.xlsx")
    wb = load_workbook(packet_dir / "R2_ReviewDefinisi.xlsx")
    wb["Review Definisi"].cell(2, 4).value = "Sesuai"
    wb.save(returned / "R2_ReviewDefinisi.xlsx")

    out = tmp_path / "data"
    info = ir.import_returns(str(returned), str(out))
    assert info["explanation_v24_rows"] == 112
    assert info["review_rows"] == 3          # one flagged row per Stage A, plus R2's separate review
    with open(out / "annotations_explanations_v24.csv", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    shown = {r["shown_output"] for r in rows if r["case_id"] == bp.code_to_case()["K01"]}
    assert shown == {"committed"}
    rep = ev.run(str(out))
    assert rep["explanations_v24"]["overall"]["accept"] == {"partly": 112}


@pytest.mark.parametrize("conclusion,category", [
    ("Tidak ada diagnosis", "no_output"),
    ("Di luar cakupan", "out_of_scope"),
    ("blas (DIDUGA)", "committed"),
    ("blas (KEMUNGKINAN (bukti parsial)); tungro (KEMUNGKINAN (bukti parsial))", "possible_only"),
])
def test_shown_category(conclusion, category):
    assert bp.shown_category(conclusion) == category
