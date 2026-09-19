import argparse
import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from openpyxl import Workbook, load_workbook

from annotation import build_packet as bp

DEFAULT_OUT = os.path.join(BASE_DIR, "annotation", "packet", "v24")
FIRST_ROUND_B = os.path.join(BASE_DIR, "annotation", "returned", "R1_TahapB.xlsx")
RULESET = "2.4.0"

STAGE_B_V24_INSTRUCTIONS = [
    "Anotasi Kasus Penyakit Padi — Tahap B (putaran kedua): Penilaian Penjelasan Sistem versi baru",
    "",
    "Sistem RiceKG telah diperbarui setelah putaran pertama: (1) gejala setiap kasus kini dimasukkan ke sistem "
    "berdasarkan kesepakatan kedua anotator pada Tahap A, dan (2) aturan diagnosis diperbarui berdasarkan literatur. "
    "Karena itu kesimpulan dan penjelasan sistem pada banyak kasus berbeda dari putaran pertama.",
    "",
    "Mohon nilai setiap kasus seolah-olah baru pertama kali melihatnya. Jangan membuka atau menyalin jawaban Tahap B "
    "putaran pertama; kami membutuhkan penilaian atas keluaran sistem yang sekarang.",
    "",
    "Bila sistem tidak menyimpulkan apa pun, penjelasannya kini menyebutkan aturan yang paling mendekati terpenuhi, "
    "gejala yang sudah teramati, dan gejala yang masih perlu diperiksa.",
    "",
] + bp.STAGE_B_INSTRUCTIONS[4:]

REVIEW_INSTRUCTIONS = [
    "Review Definisi Istilah Ontologi RiceKG",
    "",
    "Lembar 'Review Definisi' berisi definisi setiap istilah pengamatan dan penyakit yang dipakai sistem. "
    "Sebagian definisi telah direvisi berdasarkan masukan anotator lain. Mohon nilai secara mandiri apakah "
    "setiap definisi sudah sesuai dengan pengertian agronomis yang berlaku (Sesuai / Perlu revisi / Tidak sesuai), "
    "dan tuliskan usulan perbaikan bila perlu.",
    "",
    "Sel berwarna kuning adalah sel isian. Mohon tidak mengubah kolom lain, urutan baris, atau nama lembar.",
]

LETTER = """Yth. Bapak/Ibu {rater_label},

Terima kasih atas kesediaan Bapak/Ibu mengisi anotasi kasus penyakit padi pada putaran pertama.
Jawaban Tahap A Bapak/Ibu dan anotator lain telah kami gunakan: gejala setiap kasus kini dimasukkan
ke sistem berdasarkan kesepakatan kedua anotator. Selain itu, aturan diagnosis sistem kami perbarui
berdasarkan literatur. Akibatnya, kesimpulan dan penjelasan sistem pada sebagian besar kasus berubah,
sehingga penilaian Tahap B putaran pertama tidak lagi menggambarkan sistem yang sekarang.

Kami mohon kesediaan Bapak/Ibu untuk:
{tasks}

Petunjuk ada pada lembar pertama setiap berkas. Kode kasus sama dengan putaran pertama. Mohon
dikerjakan secara mandiri, tanpa berdiskusi dengan anotator lain dan tanpa membuka jawaban putaran
pertama. Perkiraan waktu: {estimate}.

Jawaban hanya dilaporkan secara agregat dengan kode anotator, dan keikutsertaan tetap bersifat
sukarela. Mohon berkas dikembalikan tanpa mengganti nama berkas.

Terima kasih atas bantuan Bapak/Ibu.

Hormat kami,
Tim RiceKG
"""


def v24_cases():
    # First-packet cases with RiceKG v2.4.0 outputs on the expert-consensus encoding.
    from ricekg import evaluate
    defs = bp.load_definitions()
    consensus = evaluate.load_symptom_encoding(evaluate.CONSENSUS_ENCODING_CSV)
    cases = bp.load_cases()
    for c in cases:
        c["symptoms"] = consensus[c["case_id"]]
    bp.explain_cases(cases, defs)
    return cases


def first_round_conclusions(path=FIRST_ROUND_B):
    wb = load_workbook(path, read_only=True)
    rows = wb["Tugas B"].iter_rows(values_only=True)
    header = next(rows)
    return {d["Kode"]: d["Kesimpulan sistem"] for d in (dict(zip(header, r)) for r in rows) if d["Kode"]}


def build_review_only(path, rater, defs):
    wb = Workbook()
    bp.instructions(wb.active, REVIEW_INSTRUCTIONS)
    wb.active.title = "Petunjuk"
    ranges = bp.list_sheet(wb, {"review": bp.REVIEW_OPTIONS})
    bp.add_review_sheet(wb, defs, ranges["review"])
    wb.properties.creator = f"RiceKG definition review ({rater})"
    wb.save(path)


def build(out_dir=DEFAULT_OUT, n_raters=2, review=("R2",), only_changed=False):
    os.makedirs(out_dir, exist_ok=True)
    defs = bp.load_definitions()
    cases = v24_cases()
    if only_changed:
        before = first_round_conclusions()
        cases = [c for c in cases if before.get(c["code"]) != c["conclusion"]]
    raters = [f"R{i}" for i in range(1, n_raters + 1)]
    for i, rater in enumerate(raters):
        bp.build_stage_b(os.path.join(out_dir, f"{rater}_TahapB_v24.xlsx"), rater,
                         bp.rater_order(cases, i), STAGE_B_V24_INSTRUCTIONS)
        tasks = [f"1. mengisi berkas {rater}_TahapB_v24.xlsx: menilai kesimpulan dan penjelasan sistem versi "
                 f"baru untuk {len(cases)} kasus;"]
        minutes = round(len(cases) * 1.5)
        if rater in review:
            build_review_only(os.path.join(out_dir, f"{rater}_ReviewDefinisi.xlsx"), rater, defs)
            tasks.append(f"2. mengisi berkas {rater}_ReviewDefinisi.xlsx: menilai {len(defs)} definisi istilah "
                         "yang belum sempat dinilai pada putaran pertama.")
            minutes += 45
        with open(os.path.join(out_dir, f"Surat_{rater}.txt"), "w", encoding="utf-8") as f:
            f.write(LETTER.format(rater_label=f"Anotator {rater}", tasks="\n".join(tasks),
                                  estimate=f"sekitar {minutes // 60} jam {minutes % 60} menit" if minutes >= 60
                                  else f"sekitar {minutes} menit"))
    return {"cases": len(cases), "raters": raters, "review": list(review), "out_dir": out_dir}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Build the ruleset v2.4.0 re-rating packet for the raters.")
    ap.add_argument("--raters", type=int, default=2)
    ap.add_argument("--review", nargs="*", default=["R2"], help="Raters who also receive the definition review.")
    ap.add_argument("--only-changed", action="store_true")
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    info = build(args.out, args.raters, tuple(args.review), args.only_changed)
    print(f"Re-rating packet written to {info['out_dir']}: {info['cases']} cases, raters {info['raters']}, "
          f"definition review for {info['review']}")
