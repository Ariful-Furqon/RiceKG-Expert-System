import argparse
import csv
import os
import random
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from annotation.redaction import redact

FIELD_CSV = os.path.join(BASE_DIR, "data", "benchmark_field.csv")
REJECTED_CSV = os.path.join(BASE_DIR, "data", "rejected_field_candidates.csv")
DEFINITIONS_CSV = os.path.join(BASE_DIR, "ontology", "term_definitions.csv")
DEFAULT_OUT = os.path.join(BASE_DIR, "annotation", "packet")

SEED = 20260919
PRACTICE_CASES = ["FIELD_37", "FIELD_39", "FIELD_40", "FIELD_50"]
N_SYMPTOM_COLUMNS = 8

DIAGNOSIS_OPTIONS = [
    ("Blas", "Rice_Blast"),
    ("Hawar daun bakteri", "Bacterial_Leaf_Blight"),
    ("Gosong palsu", "False_Smut"),
    ("Kerdil rumput", "Rice_Grassy_Stunt"),
    ("Tungro", "Rice_Tungro_Virus"),
    ("Nematoda puru akar", "Rice_Root_Nematode"),
    ("Penyakit/hama lain di luar 6 kelas ini", "Other"),
    ("Tidak dapat ditentukan dari teks", "Undetermined"),
]
CONFIDENCE_OPTIONS = ["Yakin", "Cukup yakin", "Ragu"]
ACCEPT_OPTIONS = ["Ya", "Sebagian", "Tidak"]
LIKERT_OPTIONS = ["1", "2", "3", "4", "5"]
REVIEW_OPTIONS = ["Sesuai", "Perlu revisi", "Tidak sesuai"]

GRADE_ID = {"confirmed": "TERKONFIRMASI", "suspected": "DIDUGA", "possible": "KEMUNGKINAN (bukti parsial)"}

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------

def diagnosis_label(label, key):
    return f"{label} [{key}]"


def load_definitions():
    with open(DEFINITIONS_CSV, encoding="utf-8", newline="") as f:
        return {r["term"]: r for r in csv.DictReader(f)}


# Dropdown labels used by the September 2026 packet for terms whose labels were later revised
# (ontology v2.3.0). Returned workbooks keep the labels they were issued with.
LEGACY_LABELS = {
    "kuning mengikuti tulang daun — yellowing along leaf veins": "Yellowing_Leaf_Veins",
    "muncul pada fase berbunga hingga masak susu — onset at flowering to milky stage": "Milky_Stage_Vulnerability",
}


def symptom_labels(defs, include_legacy=False):
    # Dropdown label -> term for every observation term; labels must be unique.
    out = dict(LEGACY_LABELS) if include_legacy else {}
    for term, row in defs.items():
        if row["kind"] != "observation":
            continue
        label = f"{row['label_id']} — {row['label_en']}"
        assert label not in out, f"duplicate symptom label {label}"
        out[label] = term
    return dict(sorted(out.items()))


def case_text(row):
    # Text shown to raters: English (translated if needed), redacted; original kept if translated.
    main, removed = redact(row.get("raw_symptom_text_en") or row["raw_symptom_text"])
    if row.get("raw_symptom_text_en"):
        original, removed_orig = redact(row["raw_symptom_text"])
        main = f"{main}\n\nTeks asli: {original}"
        removed += removed_orig
    return main, removed


def load_cases():
    with open(FIELD_CSV, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    rng = random.Random(SEED)
    order = rows[:]
    rng.shuffle(order)
    cases = []
    for i, row in enumerate(order, 1):
        text, removed = case_text(row)
        cases.append({
            "code": f"K{i:02d}", "case_id": row["case_id"], "split": row["split"],
            "truth": row["diagnosis"], "text": text, "removed": removed,
            "symptoms": [row[f"symptom_{j}"] for j in range(1, 7) if row[f"symptom_{j}"]],
            "original_text": row["raw_symptom_text"],
        })
    return cases


def load_practice():
    with open(REJECTED_CSV, encoding="utf-8-sig", newline="") as f:
        rows = {r["case_id"]: r for r in csv.DictReader(f)}
    return [{"code": f"L{i}", "case_id": cid, "text": case_text(rows[cid])[0]}
            for i, cid in enumerate(PRACTICE_CASES, 1)]


def code_to_case():
    return {c["code"]: c["case_id"] for c in load_cases()}


def rater_order(cases, rater_index):
    order = cases[:]
    random.Random(SEED + rater_index).shuffle(order)
    return order


# ---------------------------------------------------------------------------
# RiceKG explanations for Stage B
# ---------------------------------------------------------------------------

def render_explanation(outputs, defs, symptoms=()):
    label = lambda t: defs[t]["label_id"] if t in defs else t
    if not outputs:
        from ricekg import model
        ab = model.explain_abstention(list(symptoms))
        text = ("Sistem tidak menyimpulkan penyakit apa pun: tidak ada aturan yang terpenuhi oleh gejala "
                "yang dicatat, dan tidak ada tanda yang menunjukkan penyebab di luar cakupan.")
        for r in ab["nearest_rules"]:
            text += (f"\nAturan terdekat: {label(r['threat'])} ({r['rule_id']}, {round(100 * r['coverage'])}% gejala teramati). "
                     f"Sudah teramati: {', '.join(label(s) for s in r['matched_symptoms'])}. "
                     f"Perlu diperiksa: {', '.join(label(s) for s in r['missing_symptoms'])}.")
        if ab["unused_observations"]:
            text += f"\nTercatat tetapi tidak dipakai aturan mana pun: {', '.join(label(s) for s in ab['unused_observations'])}."
        return text
    parts = []
    for r in outputs:
        if r["grade"] == "out_of_scope":
            signs = ", ".join(label(s) for s in r["matched_symptoms"])
            if "insect" in r["threat"]:
                parts.append(f"DI LUAR CAKUPAN: tanda sesuai dengan kerusakan serangga ({signs}); serangga tidak didiagnosis oleh sistem ini.")
            else:
                parts.append(f"DI LUAR CAKUPAN: tanda tercatat ({signs}) tidak sesuai dengan penyakit mana pun dalam cakupan sistem.")
            continue
        line = f"{label(r['threat'])} — {GRADE_ID.get(r['grade'], r['grade'])}."
        if r["fired_rules"]:
            line += f" Aturan terpenuhi: {', '.join(r['fired_rules'])}."
        else:
            line += f" Tidak ada aturan yang terpenuhi; {round(100 * r['antecedent_coverage'])}% gejala aturan relaksasi teramati."
        if r["matched_symptoms"]:
            line += f" Gejala yang cocok: {', '.join(label(s) for s in r['matched_symptoms'])}."
        if r["missing_symptoms"]:
            line += f" Gejala aturan yang tidak teramati: {', '.join(label(s) for s in r['missing_symptoms'])}."
        parts.append(line)
    return "\n".join(parts)


def ricekg_conclusion(outputs, defs):
    if not outputs:
        return "Tidak ada diagnosis"
    label = lambda t: defs[t]["label_id"] if t in defs else t
    return "; ".join("Di luar cakupan" if r["grade"] == "out_of_scope"
                     else f"{label(r['threat'])} ({GRADE_ID.get(r['grade'], r['grade'])})" for r in outputs)


def shown_category(conclusion):
    # Category of the RiceKG output a rater saw, from the Stage B conclusion text.
    text = conclusion or ""
    if text.startswith("Tidak ada diagnosis"):
        return "no_output"
    if "Di luar cakupan" in text:
        return "out_of_scope"
    if GRADE_ID["confirmed"] in text or GRADE_ID["suspected"] in text:
        return "committed"
    if GRADE_ID["possible"] in text:
        return "possible_only"
    return "unknown"


def explain_cases(cases, defs):
    from ricekg import model
    for c in cases:
        outputs = model.predict_diseases(c["symptoms"], include_possible=True)
        c["conclusion"] = ricekg_conclusion(outputs, defs)
        c["explanation"] = render_explanation(outputs, defs, c["symptoms"])
        c["recorded"] = ", ".join(defs[s]["label_id"] for s in c["symptoms"]) or "(tidak ada gejala yang terpetakan)"


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------

def header(ws, titles, widths):
    ws.append(titles)
    for i, (cell, w) in enumerate(zip(ws[1], widths), 1):
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "C2"


def list_sheet(wb, lists):
    # Hidden sheet holding dropdown lists; returns name -> absolute range.
    ws = wb.create_sheet("Daftar")
    ranges = {}
    for col, (name, values) in enumerate(lists.items(), 1):
        letter = ws.cell(row=1, column=col).column_letter
        ws.cell(row=1, column=col, value=name)
        for i, v in enumerate(values, 2):
            ws.cell(row=i, column=col, value=v)
        ranges[name] = f"Daftar!${letter}$2:${letter}${len(values) + 1}"
    ws.sheet_state = "hidden"
    return ranges


def dropdown(ws, rng, cells):
    dv = DataValidation(type="list", formula1=f"={rng}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cells)


def instructions(ws, lines):
    ws.column_dimensions["A"].width = 120
    for line in lines:
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).alignment = WRAP
    ws["A1"].font = Font(bold=True, size=14)


STAGE_A_INSTRUCTIONS = [
    "Anotasi Kasus Penyakit Padi — Tahap A",
    "",
    "Terima kasih atas kesediaan Anda. Setiap baris berisi kutipan gejala dari laporan kasus lapangan yang telah diterbitkan. "
    "Nama penyakit dan patogen telah disamarkan menjadi [disamarkan]. Kerjakan secara mandiri; jangan berdiskusi dengan anotator lain "
    "dan jangan mencari sumber aslinya.",
    "",
    "Untuk setiap kasus pada lembar 'Tugas A':",
    "1. Diagnosis — pilih satu kelas yang paling didukung oleh teks. Pilih 'Penyakit/hama lain di luar 6 kelas ini' bila gejalanya "
    "menunjukkan penyebab lain, dan 'Tidak dapat ditentukan dari teks' bila informasinya tidak cukup untuk memutuskan.",
    "2. Keyakinan — seberapa yakin Anda terhadap diagnosis tersebut.",
    "3. Gejala 1–8 — pilih dari daftar SEMUA pengamatan yang benar-benar disebutkan dalam teks. Jangan menambahkan gejala yang tidak "
    "tertulis, meskipun biasanya menyertai penyakit tersebut. Gunakan lembar 'Glosarium' untuk definisi setiap istilah.",
    "4. Gejala lain — tuliskan pengamatan dalam teks yang tidak ada padanannya di daftar.",
    "5. Catatan — opsional.",
    "",
    "Mulailah dengan lembar 'Latihan' (4 kasus). Kasus latihan tidak dianalisis; kasus ini dibahas bersama koordinator sebelum "
    "Anda mengerjakan 'Tugas A'.",
    "",
    "Lembar 'Review Definisi': nilai apakah setiap definisi istilah sudah sesuai dengan pengertian agronomis yang berlaku, "
    "dan tuliskan usulan perbaikan bila perlu.",
    "",
    "Sel berwarna kuning adalah sel isian. Mohon tidak mengubah kolom lain, urutan baris, atau nama lembar.",
    "Tahap B (penilaian penjelasan sistem) akan dikirim setelah Tahap A Anda kembalikan.",
]

STAGE_B_INSTRUCTIONS = [
    "Anotasi Kasus Penyakit Padi — Tahap B: Penilaian Penjelasan Sistem",
    "",
    "Pada tahap ini Anda menilai kesimpulan dan penjelasan sistem pakar RiceKG untuk kasus yang sama dengan Tahap A (kode kasus sama). "
    "Jangan mengubah jawaban Tahap A Anda.",
    "",
    "Kolom 'Gejala yang dicatat untuk sistem' adalah pengamatan yang dimasukkan ke sistem. Kolom 'Penjelasan sistem' menunjukkan aturan "
    "yang terpenuhi, gejala yang cocok, dan gejala aturan yang tidak teramati.",
    "Tingkat kesimpulan: TERKONFIRMASI (aturan lengkap terpenuhi), DIDUGA (aturan relaksasi terpenuhi), KEMUNGKINAN (sebagian gejala "
    "aturan teramati, tidak ada aturan yang terpenuhi), DI LUAR CAKUPAN (sistem menolak mendiagnosis).",
    "",
    "Untuk setiap kasus:",
    "B1. Dapatkah kesimpulan sistem diterima, berdasarkan teks kasus? (Ya / Sebagian / Tidak). "
    "Termasuk bila sistem tidak menyimpulkan apa pun: apakah keputusan itu tepat?",
    "B2. Apakah penalaran dari gejala ke kesimpulan benar secara agronomis? (1 = sangat tidak benar, 5 = sepenuhnya benar)",
    "B3. Seberapa lengkap penjelasannya? (1 = sangat tidak lengkap, 5 = lengkap)",
    "B4. Seberapa berguna penjelasan ini bagi petugas lapangan (POPT/PPL)? (1 = tidak berguna, 5 = sangat berguna)",
    "Komentar — opsional, terutama bila Anda memberi nilai rendah.",
]


def add_review_sheet(wb, defs, review_range):
    rv = wb.create_sheet("Review Definisi")
    header(rv, ["Istilah (ID sistem)", "Label Indonesia", "Definisi", "Penilaian", "Usulan perbaikan"], [28, 32, 70, 16, 50])
    for term, row in sorted(defs.items()):
        rv.append([term, row["label_id"], row["definition"], None, None])
        rv.cell(row=rv.max_row, column=3).alignment = WRAP
        for col in (4, 5):
            rv.cell(row=rv.max_row, column=col).fill = INPUT_FILL
    dropdown(rv, review_range, f"D2:D{rv.max_row}")
    return rv


def build_stage_a(path, rater, cases, practice, defs, sym_labels):
    wb = Workbook()
    instructions(wb.active, STAGE_A_INSTRUCTIONS)
    wb.active.title = "Petunjuk"
    ranges = list_sheet(wb, {
        "diagnosis": [diagnosis_label(l, k) for l, k in DIAGNOSIS_OPTIONS],
        "keyakinan": CONFIDENCE_OPTIONS,
        "gejala": list(sym_labels),
        "review": REVIEW_OPTIONS,
    })
    titles = (["Kode", "Teks gejala", "Diagnosis", "Keyakinan"]
              + [f"Gejala {i}" for i in range(1, N_SYMPTOM_COLUMNS + 1)] + ["Gejala lain", "Catatan"])
    widths = [7, 80, 30, 13] + [30] * N_SYMPTOM_COLUMNS + [30, 30]
    for sheet_name, rows in (("Latihan", practice), ("Tugas A", cases)):
        ws = wb.create_sheet(sheet_name, 1 if sheet_name == "Latihan" else 2)
        header(ws, titles, widths)
        for c in rows:
            ws.append([c["code"], c["text"]] + [None] * (len(titles) - 2))
            r = ws.max_row
            ws.cell(row=r, column=2).alignment = WRAP
            ws.row_dimensions[r].height = 110
            for col in range(3, len(titles) + 1):
                ws.cell(row=r, column=col).fill = INPUT_FILL
        last = ws.max_row
        dropdown(ws, ranges["diagnosis"], f"C2:C{last}")
        dropdown(ws, ranges["keyakinan"], f"D2:D{last}")
        end_col = ws.cell(row=1, column=4 + N_SYMPTOM_COLUMNS).column_letter
        dropdown(ws, ranges["gejala"], f"E2:{end_col}{last}")

    gl = wb.create_sheet("Glosarium")
    header(gl, ["Istilah (ID sistem)", "Label Indonesia", "Label Inggris", "Definisi", "Catatan cakupan"], [28, 32, 32, 70, 60])
    for term, row in sorted(defs.items()):
        if row["kind"] == "observation":
            gl.append([term, row["label_id"], row["label_en"], row["definition"], row["scope_note"]])
            for col in (4, 5):
                gl.cell(row=gl.max_row, column=col).alignment = WRAP

    add_review_sheet(wb, defs, ranges["review"])

    wb.properties.creator = f"RiceKG annotation packet ({rater})"
    wb.save(path)


def build_stage_b(path, rater, cases, instruction_lines=STAGE_B_INSTRUCTIONS):
    wb = Workbook()
    instructions(wb.active, instruction_lines)
    wb.active.title = "Petunjuk"
    ranges = list_sheet(wb, {"terima": ACCEPT_OPTIONS, "skala": LIKERT_OPTIONS})
    ws = wb.create_sheet("Tugas B", 1)
    header(ws, ["Kode", "Teks gejala", "Gejala yang dicatat untuk sistem", "Kesimpulan sistem", "Penjelasan sistem",
                "B1 Diterima?", "B2 Penalaran benar (1-5)", "B3 Kelengkapan (1-5)", "B4 Kegunaan (1-5)", "Komentar"],
           [7, 60, 35, 30, 70, 13, 13, 13, 13, 35])
    for c in cases:
        ws.append([c["code"], c["text"], c["recorded"], c["conclusion"], c["explanation"], None, None, None, None, None])
        r = ws.max_row
        for col in (2, 3, 4, 5):
            ws.cell(row=r, column=col).alignment = WRAP
        for col in range(6, 11):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        ws.row_dimensions[r].height = 140
    last = ws.max_row
    dropdown(ws, ranges["terima"], f"F2:F{last}")
    dropdown(ws, ranges["skala"], f"G2:I{last}")
    wb.properties.creator = f"RiceKG annotation packet ({rater})"
    wb.save(path)


def build_key(path, cases, raters):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kunci"
    header(ws, ["Kode", "case_id", "split", "diagnosis (label benchmark)"], [8, 12, 10, 26])
    for c in sorted(cases, key=lambda c: c["code"]):
        ws.append([c["code"], c["case_id"], c["split"], c["truth"]])
    rd = wb.create_sheet("Redaksi")
    header(rd, ["Kode", "case_id", "Teks asli", "Teks yang ditampilkan", "Yang disamarkan", "Sudah dicek koordinator"],
           [8, 12, 70, 70, 40, 14])
    for c in sorted(cases, key=lambda c: c["code"]):
        rd.append([c["code"], c["case_id"], c["original_text"], c["text"], "; ".join(c["removed"]), None])
        for col in (3, 4):
            rd.cell(row=rd.max_row, column=col).alignment = WRAP
    od = wb.create_sheet("Urutan per rater")
    od.append(["Posisi"] + raters)
    orders = [[c["code"] for c in rater_order(cases, i)] for i in range(len(raters))]
    for pos, codes in enumerate(zip(*orders), 1):
        od.append([pos, *codes])
    wb.save(path)


def build(out_dir=DEFAULT_OUT, n_raters=3, with_explanations=True):
    os.makedirs(out_dir, exist_ok=True)
    defs = load_definitions()
    sym_labels = symptom_labels(defs)
    cases = load_cases()
    practice = load_practice()
    if with_explanations:
        explain_cases(cases, defs)
    raters = [f"R{i}" for i in range(1, n_raters + 1)]
    for i, rater in enumerate(raters):
        order = rater_order(cases, i)
        build_stage_a(os.path.join(out_dir, f"{rater}_TahapA.xlsx"), rater, order, practice, defs, sym_labels)
        if with_explanations:
            build_stage_b(os.path.join(out_dir, f"{rater}_TahapB.xlsx"), rater, order)
    build_key(os.path.join(out_dir, "KUNCI_koordinator_JANGAN_DIKIRIM.xlsx"), cases, raters)
    return {"cases": len(cases), "practice": len(practice), "raters": raters, "out_dir": out_dir}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Build the multi-rater annotation packet: one workbook per rater and stage, plus the coordinator key.")
    ap.add_argument("--raters", type=int, default=3)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    info = build(args.out, args.raters)
    print(f"Packet written to {info['out_dir']}: {info['cases']} cases, {info['practice']} practice cases, raters {info['raters']}")
