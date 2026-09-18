"""
annotation/import_returns.py
----------------------------
Reads the completed rater workbooks from annotation/returned/ (R1_TahapA.xlsx,
R1_TahapB.xlsx, ...) and writes the de-identified annotation tables:

  data/annotations_diagnosis.csv     case_id, annotator, diagnosis, confidence, other_signs, notes
  data/annotations_multirater.csv    case_id, annotator_1..n, notes   (input to analysis/agreement.py;
                                     only cases every rater diagnosed)
  data/annotations_symptoms.csv      case_id, annotator, term
  data/annotations_explanations.csv  case_id, annotator, accept, reasoning, completeness, usefulness, comment
  data/definition_review.csv         term, annotator, rating, suggestion

Pseudonymous codes are mapped back to case IDs with the packet's fixed seed, so the key file is
not needed. Any value that does not match a dropdown option is reported and the import stops.

    python annotation/import_returns.py [--returned DIR] [--out DIR]
"""

import argparse
import csv
import glob
import os
import re
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from openpyxl import load_workbook

from annotation import build_packet as bp

DEFAULT_RETURNED = os.path.join(BASE_DIR, "annotation", "returned")
DEFAULT_OUT = os.path.join(BASE_DIR, "data")
_KEY_RX = re.compile(r"\[(\w+)\]\s*$")
_CONFIDENCE = {"Yakin": "high", "Cukup yakin": "medium", "Ragu": "low"}
_ACCEPT = {"Ya": "yes", "Sebagian": "partly", "Tidak": "no"}
_REVIEW = {"Sesuai": "adequate", "Perlu revisi": "needs_revision", "Tidak sesuai": "inadequate"}


class ImportError_(Exception):
    pass


def _rows(ws):
    header = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in r):
            yield dict(zip(header, r))


def _text(v):
    return "" if v is None else str(v).strip()


def read_stage_a(path, rater, codes, sym_labels, errors):
    wb = load_workbook(path, read_only=True)
    diagnoses, symptoms = [], []
    for row in _rows(wb["Tugas A"]):
        code = _text(row["Kode"])
        if code not in codes:
            errors.append(f"{os.path.basename(path)}: unknown case code {code!r}")
            continue
        case_id = codes[code]
        dx = _text(row["Diagnosis"])
        if dx:
            m = _KEY_RX.search(dx)
            if not m:
                errors.append(f"{rater} {code}: diagnosis {dx!r} is not a dropdown option")
                continue
            conf = _text(row["Keyakinan"])
            if conf and conf not in _CONFIDENCE:
                errors.append(f"{rater} {code}: confidence {conf!r} is not a dropdown option")
            diagnoses.append({"case_id": case_id, "annotator": rater, "diagnosis": m.group(1),
                              "confidence": _CONFIDENCE.get(conf, ""), "other_signs": _text(row["Gejala lain"]),
                              "notes": _text(row["Catatan"])})
        seen = set()
        for i in range(1, bp.N_SYMPTOM_COLUMNS + 1):
            label = _text(row[f"Gejala {i}"])
            if not label:
                continue
            if label not in sym_labels:
                errors.append(f"{rater} {code}: symptom {label!r} is not a dropdown option")
                continue
            term = sym_labels[label]
            if term not in seen:
                seen.add(term)
                symptoms.append({"case_id": case_id, "annotator": rater, "term": term})
    reviews = []
    for row in _rows(wb["Review Definisi"]):
        rating = _text(row["Penilaian"])
        if not rating and not _text(row["Usulan perbaikan"]):
            continue
        if rating and rating not in _REVIEW:
            errors.append(f"{rater}: review rating {rating!r} for {row['Istilah (ID sistem)']} is not a dropdown option")
        reviews.append({"term": _text(row["Istilah (ID sistem)"]), "annotator": rater,
                        "rating": _REVIEW.get(rating, ""), "suggestion": _text(row["Usulan perbaikan"])})
    return diagnoses, symptoms, reviews


def read_stage_b(path, rater, codes, errors):
    wb = load_workbook(path, read_only=True)
    out = []
    for row in _rows(wb["Tugas B"]):
        code = _text(row["Kode"])
        if code not in codes:
            errors.append(f"{os.path.basename(path)}: unknown case code {code!r}")
            continue
        accept = _text(row["B1 Diterima?"])
        scores = [_text(row[k]) for k in ("B2 Penalaran benar (1-5)", "B3 Kelengkapan (1-5)", "B4 Kegunaan (1-5)")]
        if not accept and not any(scores):
            continue
        if accept and accept not in _ACCEPT:
            errors.append(f"{rater} {code}: B1 {accept!r} is not a dropdown option")
        for s in scores:
            if s and s not in bp.LIKERT_OPTIONS:
                errors.append(f"{rater} {code}: rating {s!r} is not 1-5")
        out.append({"case_id": codes[code], "annotator": rater, "accept": _ACCEPT.get(accept, ""),
                    "reasoning": scores[0], "completeness": scores[1], "usefulness": scores[2],
                    "comment": _text(row["Komentar"])})
    return out


def _write(path, rows, fields):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def import_returns(returned=DEFAULT_RETURNED, out_dir=DEFAULT_OUT):
    codes = bp.code_to_case()
    sym_labels = bp.symptom_labels(bp.load_definitions())
    errors = []
    diagnoses, symptoms, reviews, explanations = [], [], [], []
    stage_a = sorted(glob.glob(os.path.join(returned, "R*_TahapA.xlsx")))
    if not stage_a:
        raise ImportError_(f"no R*_TahapA.xlsx files in {returned}")
    raters = []
    for path in stage_a:
        rater = os.path.basename(path).split("_")[0]
        raters.append(rater)
        d, s, r = read_stage_a(path, rater, codes, sym_labels, errors)
        diagnoses += d
        symptoms += s
        reviews += r
        b_path = os.path.join(returned, f"{rater}_TahapB.xlsx")
        if os.path.exists(b_path):
            explanations += read_stage_b(b_path, rater, codes, errors)
    if errors:
        raise ImportError_("\n".join(errors))

    os.makedirs(out_dir, exist_ok=True)
    _write(os.path.join(out_dir, "annotations_diagnosis.csv"), diagnoses,
           ["case_id", "annotator", "diagnosis", "confidence", "other_signs", "notes"])
    _write(os.path.join(out_dir, "annotations_symptoms.csv"), symptoms, ["case_id", "annotator", "term"])
    _write(os.path.join(out_dir, "annotations_explanations.csv"), explanations,
           ["case_id", "annotator", "accept", "reasoning", "completeness", "usefulness", "comment"])
    _write(os.path.join(out_dir, "definition_review.csv"), reviews, ["term", "annotator", "rating", "suggestion"])

    by_case = {}
    for d in diagnoses:
        by_case.setdefault(d["case_id"], {})[d["annotator"]] = d["diagnosis"]
    wide = []
    for case_id in sorted(by_case):
        answers = by_case[case_id]
        if all(r in answers for r in raters):
            row = {"case_id": case_id, "notes": ""}
            row.update({f"annotator_{i}": answers[r] for i, r in enumerate(raters, 1)})
            wide.append(row)
    _write(os.path.join(out_dir, "annotations_multirater.csv"), wide,
           ["case_id"] + [f"annotator_{i}" for i in range(1, len(raters) + 1)] + ["notes"])
    return {"raters": raters, "diagnoses": len(diagnoses), "complete_cases": len(wide),
            "symptom_rows": len(symptoms), "explanation_rows": len(explanations), "review_rows": len(reviews)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--returned", default=DEFAULT_RETURNED)
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    try:
        info = import_returns(args.returned, args.out)
    except ImportError_ as exc:
        print(f"[IMPORT STOPPED]\n{exc}")
        sys.exit(1)
    print(info)
